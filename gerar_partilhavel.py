#!/usr/bin/env python3
"""
Ti'Piedade — Gerador de Dashboard Partilhavel
Uso: python3 gerar_partilhavel.py [BASE.xlsx] [--modo mobile|desktop]
Requer: pip install openpyxl
"""
import sys, os, re, json, gzip, base64, urllib.request, tarfile, argparse
from datetime import datetime
from collections import Counter
import warnings; warnings.filterwarnings('ignore')

ap = argparse.ArgumentParser()
ap.add_argument('excel', nargs='?')
ap.add_argument('--modo', choices=['mobile','desktop'], default='mobile')
args = ap.parse_args()

DIR = os.path.dirname(os.path.abspath(__file__))
INDEX = os.path.join(DIR, 'index.html')
if not os.path.exists(INDEX):
    sys.exit("ERRO: index.html nao encontrado na mesma pasta.")

xls = args.excel
if not xls:
    for f in sorted(os.listdir(DIR)):
        if f.lower().endswith(('.xlsx','.xlsm')) and not f.startswith(('~','gerar')):
            xls = os.path.join(DIR, f); break
if not xls or not os.path.exists(xls):
    sys.exit("ERRO: Excel nao encontrado. Exemplo: python3 gerar_partilhavel.py BASE.xlsx")

print(f"Excel: {os.path.basename(xls)} | Modo: {args.modo}")

try:
    from openpyxl import load_workbook
except ImportError:
    os.system(f'"{sys.executable}" -m pip install openpyxl -q')
    from openpyxl import load_workbook

print("A ler Excel...")
wb = load_workbook(xls, read_only=True, data_only=True)
ws = next((wb[n] for n in ['BASE','Base','base'] if n in wb.sheetnames), wb.active)
print(f"Sheet: {ws.title}")

MESES={'jan':1,'fev':2,'mar':3,'abr':4,'mai':5,'jun':6,'jul':7,'ago':8,'set':9,'out':10,'nov':11,'dez':12}
def norm(s):
    s=str(s or '').strip().lower()
    for a,b in [('ê','e'),('é','e'),('ã','a'),('â','a'),('ç','c'),('ó','o'),('á','a'),('í','i'),('ú','u')]: s=s.replace(a,b)
    return re.sub(r'[^a-z0-9]','',s)

rows_data,skipped,hdrs=[],0,None
for row in ws.iter_rows(values_only=True):
    if hdrs is None:
        hdrs={norm(c):i for i,c in enumerate(row) if c}; continue
    def g(*keys):
        for k in keys:
            i=hdrs.get(norm(k))
            if i is not None and i<len(row) and row[i] is not None: return row[i]
        return None
    try:
        ano_v=g('ano','year'); mes_v=str(g('mes','month','ms') or '').strip().lower()
        for a,b in [('ê','e'),('é','e'),('ã','a')]: mes_v=mes_v.replace(a,b)
        ano=int(float(str(ano_v))) if ano_v else 0
        if ano<2000: skipped+=1; continue
        mes=MESES.get(mes_v[:3],0)
        if not mes:
            try: mes=int(float(mes_v))
            except: skipped+=1; continue
        if not 1<=mes<=12: skipped+=1; continue
        valor=g('precoliquido','precolicuido','valor','price') or 0
        qtd=g('quantidade','qty','quantity') or 0
        rows_data.append({'ano':ano,'mes':mes,
            'entidade':str(g('entidade') or ''),
            'descEntidade':str(g('descricaodaentidade','nomentidade','descricaoentidade') or ''),
            'artigo':str(g('artigo') or ''),
            'descArtigo':str(g('descricaodoartigo','descricaoartigo','descartigo') or ''),
            'serie':str(g('serie','sr') or ''),
            'tipoDoc':str(g('tipodoc','tipo') or ''),
            'valor':float(str(valor).replace(',','.')) if valor else 0.0,
            'qtd':float(str(qtd).replace(',','.')) if qtd else 0.0})
    except: skipped+=1
wb.close()
print(f"Linhas: {len(rows_data)} carregadas, {skipped} ignoradas")
if not rows_data: sys.exit("ERRO: Nenhuma linha valida.")

ano_max=max(r['ano'] for r in rows_data)
rows_f=[r for r in rows_data if r['ano']>=ano_max-2]
ano_n,mes_n=max(Counter((r['ano'],r['mes']) for r in rows_f))
print(f"Periodo: {mes_n}/{ano_n} | Linhas: {len(rows_f)}")

print("A comprimir...")
dicts={k:{} for k in 'e de a da s t'.split()}
idxs={k:0 for k in dicts}
def intern(k,v):
    v=str(v or '')
    if v not in dicts[k]: dicts[k][v]=idxs[k];idxs[k]+=1
    return dicts[k][v]
compact=[[r['ano'],r['mes'],intern('e',r['entidade']),intern('de',r['descEntidade']),
    intern('a',r['artigo']),intern('da',r['descArtigo']),intern('s',r['serie']),
    intern('t',r['tipoDoc']),r['valor'],r['qtd']] for r in rows_f]
dicts_inv={}
for k,d in dicts.items():
    arr=['']*len(d); [arr.__setitem__(i,v) for v,i in d.items()]
    dicts_inv[k]=arr
payload={'v':2,'dicts':dicts_inv,'rows':compact,'period':{'ano':ano_n,'mes':mes_n},
    'custoFixo':{},'custoArtigo':{},'fileName':os.path.basename(xls),
    'exportedAt':datetime.now().isoformat()+'Z'}
compressed=gzip.compress(json.dumps(payload,separators=(',',':')).encode(),compresslevel=9)
b64=base64.b64encode(compressed).decode()
print(f"Dados: {len(b64)//1024}KB base64")

def get_lib(pkg,fname,sel):
    p=os.path.join(DIR,fname)
    if os.path.exists(p) and os.path.getsize(p)>10000: return open(p,encoding='utf-8').read()
    print(f"A descarregar {fname}...")
    urllib.request.urlretrieve(f'https://registry.npmjs.org/{pkg}','/tmp/lib.tgz')
    with tarfile.open('/tmp/lib.tgz') as t:
        m=next(x for x in t.getmembers() if sel(x.name))
        code=t.extractfile(m).read().decode()
    open(p,'w',encoding='utf-8').write(code); return code

chart=get_lib('chart.js/-/chart.js-4.4.1.tgz','chart.umd.min.js',
    lambda n:'dist/chart.umd.js'==n.replace('package/','') and 'map' not in n)
pako=get_lib('pako/-/pako-2.1.0.tgz','pako.min.js',lambda n:n.endswith('pako.es5.min.js'))

print("A montar HTML...")
tmpl=open(INDEX,encoding='utf-8').read()
app_s=tmpl.find('<script id="__appCode">')+len('<script id="__appCode">')
app_e=tmpl.rfind('</script>')
app_code=tmpl[app_s:app_e]
css_m=re.search(r'<style>([\s\S]*?)</style>',tmpl)
css=css_m.group(1) if css_m else ''
body_raw=tmpl[tmpl.find('<body>')+6:tmpl.find('<script id="__appCode">')]
body=re.sub(r'<style[^>]*>[\s\S]*?</style>','',body_raw,flags=re.IGNORECASE)
body=re.sub(r'<script[^>]*>[\s\S]*?</script>','',body,flags=re.IGNORECASE)

vp=('width=device-width, initial-scale=1.0, maximum-scale=1.0'
    if args.modo=='mobile' else 'width=device-width, initial-scale=1.0')
mes_nome=['Jan','Fev','Mar','Abr','Mai','Jun','Jul','Ago','Set','Out','Nov','Dez'][mes_n-1]

# UM ÚNICO SCRIPT: pako + dados + app
# Elimina qualquer problema de ordem de execução no Safari iOS
single_script=f"""{pako}
window.__shareMobile={'true' if args.modo=='mobile' else 'false'};
window.__sharedPayloadB64="{b64}";
{app_code}"""

html=f"""<!DOCTYPE html>
<html lang="pt-PT">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="{vp}">
<title>Ti'Piedade - {mes_nome} {ano_n}</title>
<script>{chart}</script>
<style>{css}</style>
</head>
<body>
{body}
<script>{single_script}</script>
</body>
</html>"""

opens=len(re.findall(r'<script',html,re.IGNORECASE))
closes=len(re.findall(r'</script',html,re.IGNORECASE))
fname=f"TiPiedade_{mes_nome}_{ano_n}_{args.modo}.html"
out=os.path.join(DIR,fname)
open(out,'w',encoding='utf-8').write(html)
print(f"\nGerado: {fname} ({os.path.getsize(out)//1024}KB)")
print(f"Scripts: {opens} abertos / {closes} fechados {'OK' if opens==closes else 'ERRO'}")
print(f"\nEnvia '{fname}' por email.")
print("No iPhone: guardar -> abrir com Safari")


# ── Upload automático para GitHub Pages ──────────────────────────────────────
def upload_github(html_path,
                  token="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",
                  owner="TIPiedade", repo="tipiedade-dashboard",
                  remote_file="dashboard.html"):
    try:
        import requests, base64
    except ImportError:
        os.system(f'"{sys.executable}" -m pip install requests -q')
        import requests, base64

    h = {"Authorization": f"token {token}", "Accept": "application/vnd.github.v3+json"}
    content = base64.b64encode(open(html_path,'rb').read()).decode()

    # verificar se já existe (para obter sha)
    r = requests.get(f"https://api.github.com/repos/{owner}/{repo}/contents/{remote_file}", headers=h)
    sha = r.json().get('sha') if r.status_code == 200 else None

    payload = {"message": f"Dashboard actualizado", "content": content}
    if sha: payload["sha"] = sha

    r2 = requests.put(f"https://api.github.com/repos/{owner}/{repo}/contents/{remote_file}",
        headers=h, json=payload)

    if r2.status_code in (200, 201):
        print(f"\n✅ Upload OK!")
        print(f"🔗 Link: https://{owner.lower()}.github.io/{repo}/{remote_file}")
        print("   (pode demorar 1-2 min a actualizar)")
    else:
        print(f"Erro upload: {r2.status_code} — {r2.text[:200]}")

# Fazer upload automaticamente após gerar
upload_github(out)
